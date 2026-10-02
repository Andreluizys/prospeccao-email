"""Envia os e-mails da planilha pelo seu e-mail da UOL (ou qualquer SMTP).

Uso:
  python enviar.py --teste seu@email.com    (manda 3 exemplos só pra você, não marca nada)
  python enviar.py                          (manda o lote do dia, respeitando o limite diário)
  python enviar.py --qtd 30 --so-aprovados  (só linhas com Status "Aprovado")

A senha é pedida na hora (ou lida da variável de ambiente SMTP_SENHA) e nunca é salva.
O Status é gravado na planilha a cada envio: se cair no meio, ninguém recebe duas vezes.
Endereços em descadastro.txt (um por linha) nunca recebem.
"""
import argparse, getpass, json, os, random, smtplib, ssl, sys, time
from datetime import datetime
from email.message import EmailMessage
from email.utils import formataddr, make_msgid
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
SERVIDORES = {
    'uolhost': ('smtps.uhserver.com', 465),    # UOL Host: e-mail com domínio próprio (voce@suaempresa.com.br)
    'uol': ('smtps.uol.com.br', 465),          # UOL Mail: @uol.com.br
    'bol': ('smtps.bol.com.br', 465),
    'gmail': ('smtp.gmail.com', 465),          # exige "senha de app"
    'outlook': ('smtp-mail.outlook.com', 587),
}

p = argparse.ArgumentParser()
p.add_argument('--xlsx', default='emails.xlsx')
p.add_argument('--config', default=str(RAIZ / 'config.json'))
p.add_argument('--qtd', type=int, help='máximo nesta rodada (padrão: o que falta do limite diário)')
p.add_argument('--so-aprovados', action='store_true', help='envia só Status "Aprovado" (senão, também "Rascunho")')
p.add_argument('--teste', metavar='EMAIL', help='manda 3 exemplos para este endereço e não marca nada')
p.add_argument('--sim', action='store_true', help='não pede confirmação')
a = p.parse_args()

cfg = json.loads(Path(a.config).read_text(encoding='utf-8'))
smtp_cfg = cfg['smtp']
remetente = smtp_cfg['remetente']
if smtp_cfg.get('host'):
    host, porta = smtp_cfg['host'], int(smtp_cfg.get('porta', 465))
else:
    host, porta = SERVIDORES[smtp_cfg.get('provedor', 'uolhost')]
limite_dia = int(smtp_cfg.get('limite_por_dia', 200))
pausa_min, pausa_max = smtp_cfg.get('intervalo_segundos', [40, 90])
descadastro_txt = cfg.get('linha_descadastro', '').strip()

bloqueados = set()
arq_bloq = RAIZ / 'descadastro.txt'
if arq_bloq.exists():
    bloqueados = {l.strip().lower() for l in arq_bloq.read_text(encoding='utf-8').splitlines() if l.strip()}

wb = openpyxl.load_workbook(a.xlsx)
ws = wb['Emails']
col = {c.value: i for i, c in enumerate(ws[1])}
hoje = datetime.now().strftime('%d/%m/%Y')
enviados_hoje = sum(1 for r in ws.iter_rows(min_row=2, values_only=True)
                    if str(r[col['Status']] or '').startswith(f'Enviado {hoje}'))

validos = ('Aprovado',) if a.so_aprovados else ('Aprovado', 'Rascunho')
fila = []
for row in ws.iter_rows(min_row=2):
    email = str(row[col['E-mail']].value or '').strip().lower()
    if row[col['Status']].value not in validos or '@' not in email:
        continue
    if email in bloqueados:
        row[col['Status']].value = 'Descadastrado'
        continue
    fila.append(row)

if a.teste:
    fila = fila[:3]
else:
    restante = max(0, limite_dia - enviados_hoje)
    fila = fila[:min(a.qtd or restante, restante)]
    print(f'Hoje já foram {enviados_hoje} de {limite_dia} permitidos.')

print(f'{len(fila)} e-mail(s) nesta rodada' + (f' (TESTE -> {a.teste})' if a.teste else ''))
if not fila:
    wb.save(a.xlsx)
    sys.exit()
if not a.teste and not a.sim and input('Confirmar envio? digite SIM: ').strip().upper() != 'SIM':
    sys.exit('Cancelado.')


def conectar(senha):
    if porta == 465:
        s = smtplib.SMTP_SSL(host, porta, context=ssl.create_default_context(), timeout=60)
    else:
        s = smtplib.SMTP(host, porta, timeout=60)
        s.starttls(context=ssl.create_default_context())
    s.login(remetente, senha)
    return s


senha = os.environ.get('SMTP_SENHA')
if not senha:
    print('(a senha não aparece enquanto digita; no Windows, cole com o botão direito do mouse)')
    senha = getpass.getpass(f'Senha de {remetente}: ').strip()
try:
    smtp = conectar(senha)
except smtplib.SMTPAuthenticationError:
    if porta == 465:  # algumas contas só aceitam 587/STARTTLS
        porta = 587
        try:
            smtp = conectar(senha)
        except smtplib.SMTPException:
            sys.exit(f'Senha recusada para {remetente}. Confira no webmail e rode de novo.')
    else:
        sys.exit(f'Senha recusada para {remetente}. Confira no webmail e rode de novo.')
print(f'Login OK em {host}:{porta}.')

enviados = 0
for i, row in enumerate(fila, 1):
    para = a.teste or row[col['E-mail']].value.strip()
    corpo = str(row[col['Corpo do e-mail']].value or '').replace('\\n', '\n').strip()
    if descadastro_txt and descadastro_txt not in corpo:
        corpo += '\n\n' + descadastro_txt
    msg = EmailMessage()
    msg['From'] = formataddr((smtp_cfg.get('nome_exibicao') or remetente, remetente))
    msg['To'] = para
    if smtp_cfg.get('copia') and not a.teste:
        msg['Cc'] = smtp_cfg['copia']
    msg['Reply-To'] = remetente
    msg['Subject'] = row[col['Assunto']].value
    msg['Message-ID'] = make_msgid(domain=remetente.split('@')[1])
    msg['List-Unsubscribe'] = f'<mailto:{remetente}?subject=remover>'
    msg.set_content(corpo)
    try:
        try:
            smtp.send_message(msg)
        except smtplib.SMTPServerDisconnected:
            smtp = conectar(senha)
            smtp.send_message(msg)
    except smtplib.SMTPException as e:
        print(f'[{i}] ERRO {para}: {e}')
        texto = str(e).lower()
        if not a.teste:
            row[col['Status']].value = f'Erro: {e}'[:200]
            wb.save(a.xlsx)
        if any(t in texto for t in ('limit', 'quota', 'too many', 'excedid', '421', '451')):
            print('O servidor parece ter limitado os envios. Parando por hoje para não bloquear a conta.')
            break
        continue
    enviados += 1
    print(f'[{i}/{len(fila)}] enviado -> {para}')
    if not a.teste:
        row[col['Status']].value = 'Enviado ' + datetime.now().strftime('%d/%m/%Y %H:%M')
        wb.save(a.xlsx)
    if i < len(fila):
        time.sleep(random.uniform(pausa_min, pausa_max))  # ritmo humano, evita cair em spam

smtp.quit()
wb.save(a.xlsx)
print(f'Pronto: {enviados} enviado(s).')
